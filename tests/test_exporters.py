"""Tests unitarios para los exportadores CSV y Excel de job0t."""

import csv
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from job0t.exporters.csv_exporter import CsvExporter
from job0t.exporters.excel_exporter import ExcelExporter
from job0t.models import Job


def _create_sample_jobs():
    return [
        Job(
            title="Senior Python Backend",
            company="Tech Corp",
            location="Córdoba",
            modality="Remoto",
            seniority="Senior",
            salary="$3000 USD",
            description="Python FastAPI Copilot",
            url="https://example.com/job/1",
            source="computrabajo",
            published_at=datetime(2026, 10, 1, 10, 30),
            categories=["Desarrollo de software"],
            ai_friendly=True,
            ai_evidence="copilot",
            freelance=False,
            remoto=True,
            remoto_evidence="modalidad declarada",
            status="Nueva",
        ),
        Job(
            title="Diseñador UX UI Freelance",
            company="Design Studio",
            location="Argentina",
            modality="Híbrido",
            seniority="N/D",
            salary=None,
            description="Diseño Figma freelance por proyecto",
            url="https://example.com/job/2",
            source="getonboard",
            published_at=datetime(2026, 10, 2, 14, 0),
            categories=["Diseño"],
            ai_friendly=False,
            freelance=True,
            freelance_evidence="freelance, por proyecto",
            remoto=False,
            status="Interesa",
        ),
    ]


def test_csv_exporter(tmp_path: Path):
    output_file = tmp_path / "test_jobs.csv"
    jobs = _create_sample_jobs()
    exporter = CsvExporter()

    res = exporter.export(jobs, output_file)
    assert res.exists()

    with open(output_file, "r", encoding="utf-8-sig") as f:
        reader = list(csv.reader(f))

    # Cabecera + 2 filas
    assert len(reader) == 3
    headers = reader[0]
    assert "Estado" in headers
    assert "Seniority" in headers
    assert "AI Friendly" in headers
    assert "Link" in headers

    row1 = reader[1]
    assert row1[0] == "Nueva"
    assert row1[1] == "Senior Python Backend"
    assert row1[5] == "Senior"  # Seniority
    assert row1[8] == "SI"  # AI Friendly
    assert row1[10] == "NO"  # Freelance
    assert row1[12] == "SI"  # Remoto

    row2 = reader[2]
    assert row2[0] == "Interesa"
    assert row2[5] == "N/D"
    assert row2[8] == "NO"
    assert row2[10] == "SI"


def test_excel_exporter(tmp_path: Path):
    output_file = tmp_path / "test_jobs.xlsx"
    jobs = _create_sample_jobs()
    exporter = ExcelExporter()

    res = exporter.export(jobs, output_file)
    assert res.exists()

    # Cargar y verificar con openpyxl
    wb = load_workbook(output_file, data_only=False)
    assert "Ofertas" in wb.sheetnames
    assert "Resumen" in wb.sheetnames

    ws_offers = wb["Ofertas"]
    assert ws_offers.max_row == 3
    assert ws_offers.max_column == 17

    # Verificar cabecera
    assert ws_offers.cell(row=1, column=1).value == "Estado"
    assert ws_offers.cell(row=1, column=6).value == "Seniority"
    assert ws_offers.cell(row=1, column=17).value == "Link"

    # Verificar datos e hipervínculo
    assert ws_offers.cell(row=2, column=6).value == "Senior"
    link_formula = ws_offers.cell(row=2, column=17).value
    assert link_formula == '=HYPERLINK("https://example.com/job/1", "Ver oferta")'

    # Verificar validación de datos (dropdown en Estado)
    assert len(ws_offers.data_validations.dataValidation) > 0
    dv = ws_offers.data_validations.dataValidation[0]
    assert "Nueva,Interesa,Postulado,Descartado" in dv.formula1

    # Verificar formato condicional en columnas
    assert len(ws_offers.conditional_formatting) >= 3

    # Verificar hoja Resumen
    ws_summary = wb["Resumen"]
    assert ws_summary.cell(row=1, column=1).value == "job0t — Resumen de Búsqueda"
    # Fila 5: Total Ofertas Únicas = 2
    assert ws_summary.cell(row=5, column=2).value == 2
