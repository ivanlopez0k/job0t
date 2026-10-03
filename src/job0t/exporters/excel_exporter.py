"""Exportador de ofertas a formato Excel (.xlsx) con openpyxl y formato profesional."""

from datetime import datetime
from pathlib import Path
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from job0t.exporters.base import BaseExporter
from job0t.models import Job, RunStats


class ExcelExporter(BaseExporter):
    """Genera un archivo Excel profesional con enlaces clickeables, dropdowns y formato condicional."""

    HEADERS = [
        "Estado",
        "Título",
        "Empresa",
        "Ubicación",
        "Modalidad",
        "Seniority",
        "Salario",
        "Categorías",
        "AI Friendly",
        "Evidencia AI",
        "Freelance",
        "Evidencia Freelance",
        "Remoto",
        "Evidencia Remoto",
        "Portal",
        "Fecha Publicación",
        "Link",
    ]

    def export(self, jobs: List[Job], output_path: Path, stats: Optional[RunStats] = None) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()

        # 1. Hoja "Ofertas"
        ws_offers = wb.active
        ws_offers.title = "Ofertas"
        self._build_offers_sheet(ws_offers, jobs)

        # 2. Hoja "Resumen"
        ws_summary = wb.create_sheet(title="Resumen")
        self._build_summary_sheet(ws_summary, jobs, stats)

        wb.save(output_path)
        return output_path

    def _build_offers_sheet(self, ws, jobs: List[Job]) -> None:
        # Paleta de estilos
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        link_font = Font(name="Calibri", size=11, color="0563C1", underline="single")
        regular_font = Font(name="Calibri", size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # Escribir encabezados
        ws.append(self.HEADERS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Escribir filas de datos
        for row_idx, job in enumerate(jobs, start=2):
            ws.row_dimensions[row_idx].height = 20
            pub_date = job.published_at.strftime("%Y-%m-%d %H:%M") if job.published_at else "N/D"

            row_data = [
                job.status,
                job.title,
                job.company,
                job.location,
                job.modality,
                job.seniority,
                job.salary or "N/D",
                "; ".join(job.categories),
                "SI" if job.ai_friendly else "NO",
                job.ai_evidence,
                "SI" if job.freelance else "NO",
                job.freelance_evidence,
                "SI" if job.remoto else "NO",
                job.remoto_evidence,
                job.source,
                pub_date,
            ]
            ws.append(row_data)

            # Celda de Enlace clickeable (Columna 17: Q)
            link_cell = ws.cell(row=row_idx, column=17)
            escaped_url = job.url.replace('"', '""')
            link_cell.value = f'=HYPERLINK("{escaped_url}", "Ver oferta")'
            link_cell.font = link_font
            link_cell.alignment = center_align

            # Aplicar bordes y alineaciones a la fila
            for col_idx in range(1, 18):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                if col_idx not in (17,):
                    c.font = regular_font
                if col_idx in (1, 6, 9, 11, 13, 15, 16, 17):
                    c.alignment = center_align
                else:
                    c.alignment = left_align

        max_row = max(len(jobs) + 1, 2)

        # Congelar panel superior
        ws.freeze_panes = "A2"

        # Activar autofiltro
        ws.auto_filter.ref = f"A1:{get_column_letter(len(self.HEADERS))}{max_row}"

        # Validación de datos: Dropdown interactivo en columna Estado (A)
        dv = DataValidation(
            type="list",
            formula1='"Nueva,Interesa,Postulado,Descartado"',
            allow_blank=True,
        )
        dv.error = "Por favor seleccione un estado de la lista."
        dv.errorTitle = "Estado inválido"
        ws.add_data_validation(dv)
        dv.add(f"A2:A{max_row}")

        # Formato condicional: Si es 'SI' pintar verde suave (#E2EFDA)
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        green_font = Font(name="Calibri", size=11, bold=True, color="375623")
        rule_green = CellIsRule(operator="equal", formula=['"SI"'], fill=green_fill, font=green_font)

        # Aplicar a AI Friendly (I), Freelance (K), Remoto (M)
        for col_letter in ["I", "K", "M"]:
            ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{max_row}", rule_green)

        # Ajuste dinámico de ancho de columnas
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                if val.startswith("="):
                    max_len = max(max_len, 12)
                else:
                    max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    def _build_summary_sheet(self, ws, jobs: List[Job], stats: Optional[RunStats]) -> None:
        title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        subtitle_font = Font(name="Calibri", size=11, bold=True, color="333333")
        regular_font = Font(name="Calibri", size=11)
        bold_font = Font(name="Calibri", size=11, bold=True)
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

        ws.cell(row=1, column=1, value="job0t — Resumen de Búsqueda").font = title_font
        ws.cell(row=2, column=1, value=f"Ejecución: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = regular_font

        # Métricas generales
        total_unique = len(jobs)
        ai_count = sum(1 for j in jobs if j.ai_friendly)
        freelance_count = sum(1 for j in jobs if j.freelance)
        remoto_count = sum(1 for j in jobs if j.remoto)

        ws.cell(row=4, column=1, value="Métrica").font = subtitle_font
        ws.cell(row=4, column=2, value="Cantidad").font = subtitle_font
        ws.cell(row=4, column=1).fill = header_fill
        ws.cell(row=4, column=2).fill = header_fill

        metrics = [
            ("Total Ofertas Únicas", total_unique),
            ("AI Friendly", ai_count),
            ("Freelance / Contrato", freelance_count),
            ("Remoto", remoto_count),
        ]

        curr_row = 5
        for label, val in metrics:
            ws.cell(row=curr_row, column=1, value=label).font = regular_font
            ws.cell(row=curr_row, column=2, value=val).font = bold_font
            curr_row += 1

        # Desglose por Portal
        curr_row += 1
        ws.cell(row=curr_row, column=1, value="Portal de Origen").font = subtitle_font
        ws.cell(row=curr_row, column=2, value="Ofertas").font = subtitle_font
        ws.cell(row=curr_row, column=1).fill = header_fill
        ws.cell(row=curr_row, column=2).fill = header_fill
        curr_row += 1

        sources_tally = {}
        for j in jobs:
            sources_tally[j.source] = sources_tally.get(j.source, 0) + 1

        for src, count in sources_tally.items():
            ws.cell(row=curr_row, column=1, value=src.capitalize()).font = regular_font
            ws.cell(row=curr_row, column=2, value=count).font = bold_font
            curr_row += 1

        # Desglose por Categoría
        curr_row += 1
        ws.cell(row=curr_row, column=1, value="Categoría").font = subtitle_font
        ws.cell(row=curr_row, column=2, value="Ofertas").font = subtitle_font
        ws.cell(row=curr_row, column=1).fill = header_fill
        ws.cell(row=curr_row, column=2).fill = header_fill
        curr_row += 1

        cat_tally = {}
        for j in jobs:
            for cat in j.categories:
                cat_tally[cat] = cat_tally.get(cat, 0) + 1

        for cat, count in cat_tally.items():
            ws.cell(row=curr_row, column=1, value=cat).font = regular_font
            ws.cell(row=curr_row, column=2, value=count).font = bold_font
            curr_row += 1

        # Desglose por Seniority
        curr_row += 1
        ws.cell(row=curr_row, column=1, value="Seniority").font = subtitle_font
        ws.cell(row=curr_row, column=2, value="Ofertas").font = subtitle_font
        ws.cell(row=curr_row, column=1).fill = header_fill
        ws.cell(row=curr_row, column=2).fill = header_fill
        curr_row += 1

        seniority_tally = {}
        for j in jobs:
            seniority_tally[j.seniority] = seniority_tally.get(j.seniority, 0) + 1

        for sen, count in sorted(seniority_tally.items(), key=lambda x: x[1], reverse=True):
            ws.cell(row=curr_row, column=1, value=sen).font = regular_font
            ws.cell(row=curr_row, column=2, value=count).font = bold_font
            curr_row += 1

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 15
